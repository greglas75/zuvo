#!/usr/bin/env python3
"""Survey translation QA — deterministic checks.

  integrity <export.xlsx> --script el              Phase 1 byte-level report
  diff <old.xlsx> <new.xlsx>                        cell-by-cell diff on aligned keys (source first)
  verify <old.xlsx> <new.xlsx> --workbook QA.xlsx   diff + classify every prior workbook row
                                                    (closed clean / partial / untouched / regressed)
  simulate <export.xlsx> --workbook QA.xlsx --out patched.xlsx
                                                    apply the workbook's non-FLAG rows to a copy of the
                                                    export (column E by key only) and re-run integrity on
                                                    the result — anything new is an error the workbook
                                                    itself would introduce (adversarial check)
Add --json for machine-readable output. Add --strict to exit 1 when the run found something that
blocks delivery (diff/verify: a source change that is not an exporter difference; simulate: new
issues, new grid mismatches, stale rows or keys missing from the export). Requires openpyxl — run
through scripts/stqa.sh, which bootstraps it.
"""
import argparse
import collections
import json
import re
import sys
import unicodedata

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font
except ModuleNotFoundError:  # a bare ImportError trace reads like a script defect; it is a setup step
    sys.exit("openpyxl is missing — run it through scripts/stqa.sh (bootstraps a venv) "
             "or: python3 -m pip install openpyxl")

import stqa_fonts as fonts

NB = "\u00a0"
SCRIPTS = {
    "el": r"[Ͱ-Ͽἀ-῿]", "cyr": r"[Ѐ-ӿ]", "ar": r"[؀-ۿ]",
    "he": r"[֐-׿]", "hy": r"[԰-֏]", "ka": r"[Ⴀ-ჿ]",
    "th": r"[฀-๿]", "ja": r"[぀-ヿ一-鿿]", "ko": r"[가-힯]",
    "zh": r"[一-鿿]", "latin": None,
}
SPECIALS = {NB: "NBSP", "\u202f": "NNBSP", "\u2009": "THIN", "\u200b": "ZWSP", "\u200c": "ZWNJ",
            "\u200d": "ZWJ", "\ufeff": "BOM", "\u2028": "LSEP", "\t": "TAB", "\r": "CR"}
QA_COLS = ["n", "sev", "cat", "gid", "ev", "ql", "src", "cur", "prop", "com"]


def s(x): return "" if x is None else str(x)


def load(path):
    """Rows of an export. Fails loudly on a file that is not an export: a corrections workbook or a
    random sheet passed by mistake would otherwise produce a confident, meaningless report."""
    ws = load_workbook(path).active
    hdr = [s(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    low = [h.strip().lower() for h in hdr]
    if len(low) < 5 or "global id" not in low[0] or "export variable" not in low[1]:
        sys.exit(f"{path}: not a translation export — expected header 'Global Id | Export Variable | "
                 f"Question Label | Original Text | Translated Text', got {hdr[:5]}")
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        r = list(r) + [None] * 7
        rows.append(dict(gid=s(r[0]), ev=s(r[1]), ql=s(r[2]), src=s(r[3]), tgt=s(r[4])))
    return hdr, rows


def load_qa(path):
    """Rows of a corrections workbook (10-column contract)."""
    ws = load_workbook(path).active
    return [dict(zip(QA_COLS, [s(x) for x in (list(r) + [None] * 10)[:10]], strict=True))
            for r in ws.iter_rows(min_row=2, values_only=True) if any(x is not None for x in r)]


def norm_html(x):
    """Cross-exporter normalisation: entities, <br> forms, newline after <br>."""
    x = x.replace("&nbsp;", NB).replace("<br />\n", "<br/>")
    return x.replace("<br />", "<br/>").replace("<br>", "<br/>")


def tags(x): return [t.replace(" ", "").replace("/>", ">") for t in re.findall(r"<[^>]+>", x)]


def is_option_or_statement(ev):
    """`…-aN` option rows and `…-sN` statement rows. Empty Export Variables exist (Translation File
    exporter: `…|0` rows) — they are neither, and must not crash the check."""
    return ev.rsplit("-", 1)[-1][:1] in ("a", "s") if ev else False


def integrity(path, script):
    hdr, rows = load(path)
    tgt_re = SCRIPTS.get(script)
    if script not in SCRIPTS:
        sys.exit(f"unknown --script {script!r}; choose one of {', '.join(SCRIPTS)}")
    rep = {"file": path, "rows": len(rows), "header": hdr,
           "empty_export_variable_rows": sum(1 for r in rows if not r["ev"])}
    for col in ("src", "tgt"):
        cnt, cells = collections.Counter(), collections.Counter()
        for r in rows:
            for ch, nm in SPECIALS.items():
                k = r[col].count(ch)
                if k: cnt[nm] += k; cells[nm] += 1
        rep[f"{col}_specials"] = {"chars": dict(cnt), "cells": dict(cells)}
        rep[f"{col}_nbsp_cells"] = [(r["ev"] or r["gid"], r[col].count(NB)) for r in rows if NB in r[col]]
        rep[f"{col}_entities"] = dict(collections.Counter(
            e for r in rows for e in re.findall(r"&[a-zA-Z#0-9]+;", r[col])))
        rep[f"{col}_lead_trail"] = [r["ev"] or r["gid"] for r in rows if r[col] != r[col].strip()]
        rep[f"{col}_double_space"] = [r["ev"] or r["gid"] for r in rows
                                     if "  " in r[col] and "<style" not in r[col]]
        rep[f"{col}_non_nfc"] = [r["ev"] or r["gid"] for r in rows
                                 if unicodedata.normalize("NFC", r[col]) != r[col]]
    issues = []
    latin_tokens = collections.Counter()
    for r in rows:
        t, name = r["tgt"], (r["ev"] or r["gid"])
        if not t:
            if r["src"].strip() and not r["ev"].endswith("-e1"): issues.append((name, "EMPTY target"))
            continue
        if re.search(r"[Ѐ-ӿ]", t) and script != "cyr": issues.append((name, "Cyrillic in target"))
        for tok in re.findall(r"\S+", re.sub(r"<[^>]+>", " ", t)):
            if tgt_re and re.search(tgt_re, tok) and re.search(r"[A-Za-z]", tok):
                issues.append((name, f"mixed-script token {tok!r}"))
        if "<style" not in t:
            for tok in re.findall(r"[A-Za-z][A-Za-z0-9.]*", re.sub(r"<[^>]+>", " ", t)):
                latin_tokens[tok] += 1
            st, tt = tags(norm_html(r["src"])), tags(norm_html(t))
            if [x for x in st if x != "<br>"] != [x for x in tt if x != "<br>"]:
                issues.append((name, f"TAG parity {st} -> {tt}"))
            if re.findall(r"\{\{.*?\}\}", r["src"]) != re.findall(r"\{\{.*?\}\}", t):
                issues.append((name, "PLACEHOLDER parity"))
            num_src, num_tgt = re.findall(r"\d+", norm_html(r["src"])), re.findall(r"\d+", t)
            if num_src != num_tgt: issues.append((name, f"NUMBER parity {num_src} -> {num_tgt}"))
        else:
            a, b = norm_html(r["src"]), norm_html(t)
            body_a = a.split("<style", 1)[1] if "<style" in a else ""
            body_b = b.split("<style", 1)[1]
            if body_a != body_b:
                ws_eq = re.sub(r"\s+", " ", body_a) == re.sub(r"\s+", " ", body_b)
                issues.append((name, "STYLE block differs from source"
                               + (" (whitespace only)" if ws_eq else "")))
            if re.search(r"<style.*?<br", b, re.S): issues.append((name, "<br> INSIDE <style> — CSS is dead"))
            if re.search(r"<style.*?&[a-z#0-9]+;", b, re.S):
                issues.append((name, "entity INSIDE <style> — CSS is dead"))
        if (t.strip() == r["src"].strip() and "<style" not in t
                and not re.fullmatch(r"[\d\s%€$–\-,.<>/br ]+", t.strip())):
            issues.append((name, "target == source"))
    rep["issues"] = issues
    rep["tgt_latin_tokens"] = latin_tokens.most_common()
    # identity groups: identical source strings must have identical targets
    groups = collections.defaultdict(list)
    for r in rows:
        if r["src"].strip() and is_option_or_statement(r["ev"]): groups[norm_html(r["src"]).strip()].append(r)
    rep["grid_mismatches"] = []
    for src, rs in groups.items():
        vals = {norm_html(r["tgt"]).replace("<br/>", "") for r in rs}
        if len(rs) > 1 and len(vals) > 1:
            rep["grid_mismatches"].append(
                {"source": src[:70], "cells": [(r["ev"], r["tgt"][:80]) for r in rs],
                 "nbsp_only": len({v.replace(NB, " ") for v in vals}) == 1})
    return rep


def diff(old_path, new_path):
    _, a = load(old_path); _, b = load(new_path)
    ka = [(r["gid"], r["ev"]) for r in a]; kb = [(r["gid"], r["ev"]) for r in b]
    dup = [k for k, n in collections.Counter(kb).items() if n > 1]
    A = {(r["gid"], r["ev"]): r for r in a}; B = {(r["gid"], r["ev"]): r for r in b}
    out = {"only_old": sorted(set(ka) - set(kb)), "only_new": sorted(set(kb) - set(ka)),
           "duplicate_keys_in_new": dup, "source_changes": [], "target_changes": []}
    for k in kb:
        if k not in A: continue
        if A[k]["src"] != B[k]["src"]:
            out["source_changes"].append(
                {"gid": k[0], "ev": k[1], "old": A[k]["src"], "new": B[k]["src"],
                 "same_after_normalisation": norm_html(A[k]["src"]) == norm_html(B[k]["src"])})
        if A[k]["tgt"] != B[k]["tgt"]:
            out["target_changes"].append({"gid": k[0], "ev": k[1], "old": A[k]["tgt"], "new": B[k]["tgt"]})
    return out, A, B


def _locate(B, by_ev, gid, ev):
    """A workbook row's cell in the new export: by (Global Id, EV); EV alone only when the workbook
    carries no Global Id AND the EV is unique — empty and repeated EVs exist and must not be merged."""
    if (gid, ev) in B: return (gid, ev)
    if not gid and ev and len(by_ev.get(ev, [])) == 1: return by_ev[ev][0]
    return None


def verify(old_path, new_path, wb_path):
    d, A, B = diff(old_path, new_path)
    by_ev = collections.defaultdict(list)
    for k in B: by_ev[k[1]].append(k)
    classes, mapped = [], set()
    for q in load_qa(wb_path):
        k = _locate(B, by_ev, q["gid"], q["ev"])
        if k is None: classes.append((q["n"], q["sev"], q["ev"], "KEY MISSING")); continue
        r = B[k]
        if q["com"].startswith("SOURCE→"):
            want = q["com"].split("\n")[0][len("SOURCE→"):].strip()
            st = "closed clean" if r["src"] == want or norm_html(r["src"]) == norm_html(want) else "untouched"
        elif q["sev"] == "FLAG" or not q["prop"]:
            st = "flag (no target action)"
        elif r["tgt"] == q["prop"]: st = "closed clean"
        elif r["tgt"] == q["cur"]: st = "untouched"
        else: st = "changed — not equal to proposal (partial / acceptable / regressed: judge)"
        classes.append((q["n"], q["sev"], q["ev"], st)); mapped.add(k)
    unsolicited = [c["ev"] or c["gid"] for c in d["target_changes"] if (c["gid"], c["ev"]) not in mapped]
    reg = []  # regression hunt on changed cells
    for c in d["target_changes"]:
        t, o = c["new"], c["old"]; iss = []
        src = B[(c["gid"], c["ev"])]["src"]; is_css = "<style" in t
        if t.count(NB) > o.count(NB): iss.append(f"NBSP +{t.count(NB) - o.count(NB)}")
        if t != t.strip() and src == src.strip(): iss.append("lead/trail ws")
        if not is_css and re.search(r"\s[,.;:]", t): iss.append("space before punct")
        if unicodedata.normalize("NFC", t) != t: iss.append("non-NFC")
        if re.search(r"&[a-z#0-9]+;", t) and not re.search(r"&[a-z#0-9]+;", o):
            iss.append("entity introduced")
        if ("“" in t or "”" in t) and not ("“" in o or "”" in o): iss.append("curly quotes introduced")
        if iss: reg.append((c["ev"] or c["gid"], iss))
    nb_old = sum(r["tgt"].count(NB) for r in A.values()); nb_new = sum(r["tgt"].count(NB) for r in B.values())
    cells_old = sum(NB in r["tgt"] for r in A.values()); cells_new = sum(NB in r["tgt"] for r in B.values())
    return {"diff": d, "classes": classes, "unsolicited_changed_cells": unsolicited,
            "regressions_in_changed_cells": reg,
            "nbsp_delta": {"chars": (nb_old, nb_new), "cells": (cells_old, cells_new)},
            "fix_rate": dict(collections.Counter(f"{c[1]}: {c[3]}" for c in classes))}


def simulate(export_path, wb_path, out_path, script, font=None):
    """Apply non-FLAG rows (col I non-empty, != col H) to column E by (Global Id, EV). Never touches col D."""
    wb = load_workbook(export_path); ws = wb.active
    todo = {(q["gid"], q["ev"]): q for q in load_qa(wb_path)
            if q["sev"] != "FLAG" and q["prop"] and q["prop"] != q["cur"]}
    name, size, note = fonts.resolve(script, texts=[q["prop"] for q in todo.values()], font=font)
    applied, stale, missing = [], [], set(todo)
    for row in ws.iter_rows(min_row=2):
        k = (s(row[0].value), s(row[1].value))
        if k in todo:
            q = todo[k]; missing.discard(k)
            if s(row[4].value) != q["cur"]:
                stale.append((k[1] or k[0], "Current Translation in the workbook != the export cell "
                                            "— the workbook was built on another version"))
            row[4].value = q["prop"]
            # the patched copy is read by a human too: a cell the client's export styled in a
            # Latin-only font would show the applied target text as boxes
            row[4].font = Font(name=name, size=row[4].font.size or size, bold=row[4].font.bold)
            applied.append(k[1] or k[0])
    wb.save(out_path)
    print(f"font (patched cells): {note}", file=sys.stderr)
    before = integrity(export_path, script); after = integrity(out_path, script)
    new_issues = sorted(set(map(tuple, after["issues"])) - set(map(tuple, before["issues"])))
    new_grid = [g for g in after["grid_mismatches"] if g not in before["grid_mismatches"]]
    return {"applied": applied, "stale_rows": stale, "keys_not_in_export": sorted(missing),
            "nbsp_cells_before_after": (len(before["tgt_nbsp_cells"]), len(after["tgt_nbsp_cells"])),
            "issues_before_after": (len(before["issues"]), len(after["issues"])),
            "NEW_issues_introduced_by_workbook": new_issues, "NEW_grid_mismatches": new_grid,
            "grid_mismatches_remaining": after["grid_mismatches"], "patched_file": out_path}


def blocking(cmd, out):
    """What --strict treats as a failed run, per command."""
    if cmd in ("diff", "verify"):
        d = out["diff"] if cmd == "verify" else out
        return [f"source changed: {c['ev'] or c['gid']}" for c in d["source_changes"]
                if not c["same_after_normalisation"]]
    if cmd == "simulate":
        return ([f"new issue: {i}" for i in out["NEW_issues_introduced_by_workbook"]]
                + [f"new grid mismatch: {g['source']}" for g in out["NEW_grid_mismatches"]]
                + [f"stale row: {x[0]}" for x in out["stale_rows"]]
                + [f"key not in export: {k}" for k in out["keys_not_in_export"]])
    return []


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--json", action="store_true"); common.add_argument("--strict", action="store_true")
    p = sub.add_parser("integrity", parents=[common])
    p.add_argument("export"); p.add_argument("--script", default="latin")
    p = sub.add_parser("diff", parents=[common]); p.add_argument("old"); p.add_argument("new")
    p = sub.add_parser("simulate", parents=[common])
    p.add_argument("export"); p.add_argument("--workbook", required=True)
    p.add_argument("--out", required=True); p.add_argument("--script", default="latin")
    p.add_argument("--font", default=None)
    p = sub.add_parser("verify", parents=[common])
    p.add_argument("old"); p.add_argument("new"); p.add_argument("--workbook", required=True)
    p.add_argument("--script", default="latin")
    a = ap.parse_args()
    if a.cmd == "integrity": out = integrity(a.export, a.script)
    elif a.cmd == "diff": out = diff(a.old, a.new)[0]
    elif a.cmd == "simulate": out = simulate(a.export, a.workbook, a.out, a.script, a.font)
    else: out = verify(a.old, a.new, a.workbook)
    if a.json: print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    else:
        for k, v in out.items():
            if isinstance(v, list) and v and isinstance(v[0], dict):
                print(f"== {k} ({len(v)})")
                for item in v:
                    print("  ", {kk: (vv[:120] if isinstance(vv, str) else vv)
                                 for kk, vv in item.items()})
            elif isinstance(v, list):
                print(f"== {k} ({len(v)})"); [print("  ", x) for x in v[:200]]
            else: print(f"== {k}: {v}")
    if a.strict:
        bad = blocking(a.cmd, out)
        if bad:
            print(f"STRICT: {len(bad)} blocking finding(s)", file=sys.stderr)
            for b in bad[:20]: print("  ", b, file=sys.stderr)
            sys.exit(1)


if __name__ == "__main__":
    main()
