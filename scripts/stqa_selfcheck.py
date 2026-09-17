#!/usr/bin/env python3
"""Self-check of the skill's scripts on synthetic exports. Run after installing or editing them:

    stqa.sh selfcheck

Builds tiny exports in a temp dir and asserts the behaviours each rule in SKILL.md depends on,
including the ones that are invisible until a client opens the file: the font of every cell, the
workbook's default font, and the right-to-left sheet flag. Prints one line per check; exit 0 = all held.
"""
import json
import os
import subprocess
import sys
import tempfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import stqa_fonts as fonts
import stqa_checks as qc
from stqa_workbook import NBSP_MARK, QABook
from stqa_adversarial import AdvBook, load_primary
from stqa_reconcile import parse_args, reconcile

NB = "\u00a0"
HDR = ["Global Id", "Export Variable", "Question Label", "Original Text", "Translated Text",
       "Original Image Title", "Translated Image Title"]
STYLE_SRC = "Pick one<style>\n.a{color:red}\n</style>"


def export(path, rows):
    wb = Workbook(); ws = wb.active; ws.append(HDR)
    for r in rows: ws.append(list(r) + [None, None])
    wb.save(path)


def base_rows():
    return [
        ("1|10", "p1-q1", "S1", "Where do you live?", "Που κατοικείτε;"),
        ("1|10|1", "p1-q1-a1", "", "Other (please state)", "Άλλο (διευκρινίστε)"),
        # grid identity: the same source string, target differs by an NBSP only
        ("1|11|1", "p1-q2-a1", "", "Other (please state)", f"Άλλο{NB}(διευκρινίστε)"),
        ("1|12", "p1-q3", "Q3", "Spend up to €1,000 on {{brand}}", "Δαπάνη έως €1,000 σε {{brand}}"),
        ("1|13", "p1-q4", "Q4", STYLE_SRC, "Επιλέξτε ένα<style><br />\n.a{color:red}<br />\n</style>"),
        ("1|14|0", "", "", "<br />", ""),                       # empty EV (Translation File exporter)
        ("1|15|0", "", "", "Hidden note", ""),                                              # second empty EV
        ("1|16", "p1-q5", "Q5", "Play hall", "Play hall"),
    ]


passed = 0


def ok(name, cond, detail=""):
    global passed
    if not cond: print(f"FAIL  {name}  {detail}"); sys.exit(1)
    passed += 1; print(f"ok    {name}")


def refuses(fn, kind=AssertionError, needle=""):
    """True when fn() is refused by `kind`, optionally with `needle` in the message."""
    try: fn()
    except kind as e: return needle in str(e)
    except SystemExit as e: return kind is SystemExit and needle in str(e)
    return False


def _check(book, sev, cat, addr, proposed, gloss):
    """add() then the contract check — the pair a caller always runs together."""
    book.add(sev, cat, addr, proposed=proposed, gloss=gloss); book.check()


def _add_no_gloss(exp, primary):
    a = AdvBook(exp, primary, script="el")
    for i in (1, 2, 3, 4): a.confirm(i, "Reproduced against the source and the termbase.")
    a.add("HIGH", "Untranslated", "p1-q5", "Left in English in the target cell of a live question.",
          proposed="Αίθουσα παιχνιδιών")
    a.check()


def cells(path):
    wb = load_workbook(path)
    return wb, [c for ws in wb.worksheets for row in ws.iter_rows() for c in row]


with tempfile.TemporaryDirectory() as d:
    e1 = os.path.join(d, "r1.xlsx"); export(e1, base_rows())

    # ── Phase 1 integrity ───────────────────────────────────────────────────────────────────
    rep = qc.integrity(e1, "el")
    issues = dict((a, b) for a, b in rep["issues"])
    ok("integrity survives empty Export Variables", rep["empty_export_variable_rows"] == 2)
    ok("NBSP counted in chars and cells",
       rep["tgt_specials"]["chars"].get("NBSP") == 1 and rep["tgt_specials"]["cells"].get("NBSP") == 1)
    ok("<br> inside <style> is reported as dead CSS",
       any("INSIDE <style>" in b for a, b in rep["issues"] if a == "p1-q4"))
    ok("untranslated cell reported (target == source)", issues.get("p1-q5") == "target == source")
    ok("empty target with source text reported, keyed by Global Id",
       ("1|15|0", "EMPTY target") in [tuple(x) for x in rep["issues"]])
    ok("grid mismatch found and classified NBSP-only",
       len(rep["grid_mismatches"]) == 1 and rep["grid_mismatches"][0]["nbsp_only"])
    ok("unknown --script is refused", refuses(lambda: qc.integrity(e1, "xx"), SystemExit, "unknown --script"))

    # ── Phase 4 corrections workbook ────────────────────────────────────────────────────────
    qa = QABook(e1, script="el")
    qa.add("HIGH", "Orthography", "p1-q1", proposed="Πού κατοικείτε;",
           gloss='EN: "where (relative)" → "Where?"', comment="Interrogative needs the accent.")
    qa.add("MEDIUM", "Grid identity", "p1-q2-a1", repl=[(f"Άλλο{NB}(", "Άλλο (")],
           gloss="EN: same text; NBSP → space")
    qa.add("HIGH", "Number format", "p1-q3", repl=[("€1,000", "1.000 €")],
           gloss='EN: "€1,000" → "1.000 €" (comma is the decimal mark in el)')
    qa.flag("Design", "p1-q4", "",
            "DESIGN: layout CSS lives in translatable text — move it to a hidden-CSS question.")
    wbp = os.path.join(d, "QA.xlsx"); counts, n = qa.write(wbp)
    ok("workbook written, sorted and counted", n == 4 and counts == {"HIGH": 2, "MEDIUM": 1, "FLAG": 1})
    ok("NBSP shown with an ASCII marker, not U+237D",
       NBSP_MARK == "{NBSP}" and any(NBSP_MARK in str(c.value) for _, cs in [cells(wbp)] for c in cs))

    ok("prose in Proposed Correction is refused",
       refuses(lambda: _check(QABook(e1, script="el"), "HIGH", "x", "p1-q5",
                              "see comment — revert to source", "EN: x"), AssertionError, "meta-vocabulary"))
    ok("value change without an EN gloss is refused",
       refuses(lambda: _check(QABook(e1, script="el"), "HIGH", "x", "p1-q1", "Πού κατοικείτε;", ""),
               AssertionError, "EN gloss"))
    ok("empty Export Variable must be addressed as <Global Id>::",
       refuses(lambda: QABook(e1, script="el").row(""), KeyError)
       and QABook(e1, script="el").source("1|15|0::") == "Hidden note")
    ok("a non-export file is refused with a clear message",
       refuses(lambda: QABook(wbp, script="el"), SystemExit, "not a translation export"))

    # ── fonts: the part a client sees first ─────────────────────────────────────────────────
    name, size, note = fonts.resolve("el", texts=["Πού κατοικείτε;"])
    wb, cs = cells(wbp)
    ok(f"every cell of the Greek workbook is {name} {size}pt",
       all(c.font.name == name and c.font.size == size for c in cs if c.value is not None),
       str({c.coordinate: (c.font.name, c.font.size) for c in cs
            if c.value is not None and c.font.name != name}))
    normal = [st for st in wb._named_styles if getattr(st, "name", "") == "Normal"]
    ok("the workbook default style is that font too (a row the client adds is not Calibri)",
       normal and normal[0].font.name == name)
    ok("Greek resolves to a font that ships with Office", name == "Arial")
    ok("no cell carries the U+237D box glyph", not any("\u237d" in str(c.value) for c in cs))

    th = os.path.join(d, "th.xlsx")
    export(th, [("1|10", "p1-q1", "S1", "Postcode", "รหัสไปรษณีย์"),
                ("1|11", "p1-q2", "S2", "Play hall", "Play hall")])
    tq = QABook(th, script="th")
    tq.add("HIGH", "Untranslated", "p1-q2", proposed="ห้องเล่นเกม", gloss='EN: "Play hall" → "gaming room"')
    thp = os.path.join(d, "QA_TH.xlsx"); tq.write(thp)
    _, tcs = cells(thp)
    thai_font = {c.font.name for c in tcs if c.value is not None}
    ok("Thai does NOT get Arial (Arial has no Thai glyphs)", thai_font == {"Tahoma"}, str(thai_font))
    ok("Thai is written at 11pt, not 10", {c.font.size for c in tcs if c.value is not None} == {11.0})
    ok("the chosen Thai font really covers the text",
       not fonts.coverage("Tahoma").gaps("รหัสไปรษณีย์ห้องเล่นเกม"))
    ok("Arial is measurably wrong for Thai", bool(fonts.coverage("Arial").gaps("รหัสไปรษณีย์")))

    ko = os.path.join(d, "ko.xlsx")
    export(ko, [("1|10", "p1-q1", "S1", "Postcode", "우편번호")])
    kname, ksize, _ = fonts.resolve("ko", texts=["우편번호"])
    ok("Korean resolves to a font with Hangul",
       not fonts.coverage(kname).gaps("우편번호") if fonts.coverage(kname) else True)

    saved = fonts.FONTS["latin"]
    try:
        fonts.FONTS["latin"] = ("Arial",)                       # installed here and missing U+237D
        ok("a text no installed candidate can render is refused, not shipped as boxes",
           refuses(lambda: fonts.resolve("latin", texts=["a \u237d b"]),
                   fonts.FontError, "no font for script"))
    finally:
        fonts.FONTS["latin"] = saved
    ok("a house font passed by the human overrules the table",
       fonts.resolve("el", texts=["Πού"], font="Times New Roman")[0] == "Times New Roman")

    wb2 = load_workbook(wbp); wb2.active["C2"].font = Font(name="Calibri", size=10)
    ok("a single stray Calibri cell blocks the save",
       refuses(lambda: fonts.assert_workbook(wb2, name, size), AssertionError, "not in"))

    ar = os.path.join(d, "ar.xlsx")
    export(ar, [("1|10", "p1-q1", "S1", "Where do you live?", "أين تسكن؟")])
    aq = QABook(ar, script="ar")
    aq.add("LOW", "Punctuation", "p1-q1", proposed="أين تسكن ؟", gloss='EN: same text; spacing before "؟"')
    arp = os.path.join(d, "QA_AR.xlsx"); aq.write(arp)
    awb = load_workbook(arp)
    ok("an RTL target gets a right-to-left sheet, not only right-aligned cells",
       awb.active.sheet_view.rightToLeft is True)

    # ── simulate / verify ───────────────────────────────────────────────────────────────────
    sim = qc.simulate(e1, wbp, os.path.join(d, "patched.xlsx"), "el")
    ok("simulate applies by key and introduces nothing",
       sorted(sim["applied"]) == ["p1-q1", "p1-q2-a1", "p1-q3"]
       and not sim["NEW_issues_introduced_by_workbook"] and not sim["stale_rows"])
    ok("simulate closes the NBSP-only grid mismatch", not sim["grid_mismatches_remaining"])
    ok("--strict passes a clean simulate", qc.blocking("simulate", sim) == [])
    patched = load_workbook(os.path.join(d, "patched.xlsx")).active
    ok("simulate styles the cells it rewrites in the target-script font",
       patched["E2"].font.name == name, patched["E2"].font.name)

    e_other = os.path.join(d, "other.xlsx"); rows = base_rows(); rows[0] = rows[0][:4] + ("Πού μένετε;",)
    export(e_other, rows)
    sim2 = qc.simulate(e_other, wbp, os.path.join(d, "patched2.xlsx"), "el")
    ok("simulate reports a workbook built on another version (stale row)",
       [x[0] for x in sim2["stale_rows"]] == ["p1-q1"] and qc.blocking("simulate", sim2))

    rows = base_rows()
    rows[0] = rows[0][:4] + ("Πού κατοικείτε;",)                     # closed clean
    rows[3] = rows[3][:4] + ("Δαπάνη έως 1000 € σε {{brand}}",)      # changed, not equal to the proposal
    # source overwritten by the translation round → blocking
    rows[1] = ("1|10|1", "p1-q1-a1", "", "Other (please  state)", rows[1][4])
    rows[5] = ("1|14|0", "", "", "<br/>", "")                         # exporter-only difference
    e2 = os.path.join(d, "r2.xlsx"); export(e2, rows)
    v = qc.verify(e1, e2, wbp)
    cls = {c[2]: c[3] for c in v["classes"]}
    ok("verify: closed clean / untouched / changed classified",
       cls["p1-q1"] == "closed clean" and cls["p1-q2-a1"] == "untouched"
       and cls["p1-q3"].startswith("changed"))
    src = {c["ev"] or c["gid"]: c["same_after_normalisation"] for c in v["diff"]["source_changes"]}
    ok("verify: an exporter-only source difference is not a master change",
       src == {"p1-q1-a1": False, "1|14|0": True})
    ok("--strict blocks on a real source change only",
       qc.blocking("verify", v) == ["source changed: p1-q1-a1"])

    # ── Phase 4b adversarial file ───────────────────────────────────────────────────────────
    adv = AdvBook(e1, wbp, script="el", primary_model="GPT Astra", adversary_model="Fable 5.1 Max")
    adv.confirm(1, 'Source is a direct question; "Πού" carries the accent as an interrogative.')
    ok("a verdict for a row the primary does not have is refused",
       refuses(lambda: adv.confirm(99, "x" * 30), KeyError, "no row #99"))
    ok("the export passed instead of the workbook is refused",
       refuses(lambda: AdvBook(e1, e1, script="el"), SystemExit, "not a corrections workbook"))

    same = AdvBook(e1, wbp, script="el")
    same.dispute(2, "Repeats the primary without new evidence about the cell.",
                 proposed="Δαπάνη έως 1.000 € σε {{brand}}")
    for i in (1, 3, 4): same.confirm(i, "Reproduced against the source and the termbase.")
    ok("a DISPUTE that changes neither value nor severity is refused",
       refuses(same.check, AssertionError, "repeats the primary"))

    hedge = AdvBook(e1, wbp, script="el")
    for i in (1, 2, 3, 4): hedge.confirm(i, "Reproduced against the source and the termbase.")
    hedge.rows[0]["arg"] = "This maybe reads better, arguably, in the target language."
    ok("hedging outside a quotation is refused", refuses(hedge.check, AssertionError, "hedging"))

    adv.dispute(2, 'HIGH stands; the value is wrong: a brand placeholder takes the article '
                   '— "στο {{brand}}", '
                   "as the other brand cells of this export do.",
                proposed="Δαπάνη έως 1.000 € στο {{brand}}", gloss='EN: "on {{brand}}" → "on the {{brand}}"')
    ok("an adversarial file that leaves a primary row silent is refused",
       refuses(adv.check, AssertionError, "without a verdict"))
    adv.confirm(3, "simulate: the grid mismatch closes and the NBSP count after apply is 0.")
    adv.drop(4, "Contract: the DESIGN note duplicates the chat verdict and names no cell-level action.")
    adv.add("HIGH", "Untranslated", "p1-q5", 'Source "Play hall" is left in English in the target cell.',
            proposed="Αίθουσα παιχνιδιών", gloss='EN: "Play hall" → "games room"')
    adv.verdict_dispute("GO", "Verdict claims 0 untranslated cells; integrity reports "
                              "p1-q5 target == source.")
    advp = os.path.join(d, "ADV.xlsx"); acounts, an = adv.write(advp)
    ok("adversarial file written: one verdict per primary row + ADD + V",
       an == 6 and acounts == {"CONFIRM": 2, "DISPUTE": 2, "DROP": 1, "ADD": 1})
    _, acs = cells(advp)
    ok("the adversarial file is in the same verified font",
       all(c.font.name == name and c.font.size == size for c in acs if c.value is not None))
    ok("an ADD that changes a value without an EN gloss is refused",
       refuses(lambda: _add_no_gloss(e1, wbp), AssertionError, "EN gloss"))

    # ── reconciliation ──────────────────────────────────────────────────────────────────────
    dec = os.path.join(d, "decisions.json")
    # after write() the file is sorted DROP → DISPUTE(V) → DISPUTE(#2) → ADD → CONFIRM, and the
    # decisions are keyed by the adversarial file's own row numbers
    order = [r["verdict"] + (":" + r["ref"] if r["ref"] else "") for r in adv.rows]
    ok("adversarial rows are sorted DROP → DISPUTE → ADD → CONFIRM",
       order == ["DROP:4", "DISPUTE:V", "DISPUTE:2", "ADD", "CONFIRM:1", "CONFIRM:3"], str(order))
    with open(dec, "w") as fh:
        json.dump({"1": {"decision": "accept", "reason": "no cell-level action in the note"},
                   "2": {"decision": "accept",
                         "reason": "integrity reports target == source, so GO is unsupported"},
                   "3": {"decision": "accept", "reason": "the article is required after this preposition"},
                   "4": {"decision": "accept", "reason": "integrity reports target == source"}}, fh)
    out = os.path.join(d, "QA_reconciled.xlsx")
    res = reconcile(parse_args(["--export", e1, "--primary", wbp, "--adversarial", advp,
                                "--decisions", dec, "--out", out, "--script", "el",
                                "--adversary", "Fable 5.1 Max"]))
    got = load_primary(out)
    ok("reconcile: accepted DROP removes the row, accepted ADD adds one",
       res["rows"] == 4 and len(res["dropped"]) == 1 and len(res["added"]) == 1)
    ok("reconcile: an accepted DISPUTE value lands in Proposed Correction",
       any(q["prop"] == "Δαπάνη έως 1.000 € στο {{brand}}" for q in got.values()))
    ok("reconcile: every surviving row records the adversary's verdict",
       all("ADV (Fable 5.1 Max)" in q["com"] for q in got.values()))
    ok("reconcile: the verdict claim is reported with its decision",
       [x[0] for x in res["verdict_claims"]] == ["accept"])
    _, rcs = cells(out)
    ok("the reconciled workbook is in the verified font as well",
       all(c.font.name == name and c.font.size == size for c in rcs if c.value is not None))

    with open(dec) as fh: full = json.load(fh)
    short = dict(full); short.pop("3")
    dec2 = os.path.join(d, "d2.json")
    with open(dec2, "w") as fh: json.dump(short, fh)
    ok("reconcile refuses a non-CONFIRM row with no decision",
       refuses(lambda: reconcile(parse_args(["--export", e1, "--primary", wbp, "--adversarial", advp,
                                             "--decisions", dec2, "--out", out, "--script", "el"])),
               SystemExit, "without a decision"))
    typo = dict(full); typo["77"] = {"decision": "accept", "reason": "typo"}
    dec3 = os.path.join(d, "d3.json")
    with open(dec3, "w") as fh: json.dump(typo, fh)
    ok("reconcile refuses a decision that names no adversarial row",
       refuses(lambda: reconcile(parse_args(["--export", e1, "--primary", wbp, "--adversarial", advp,
                                             "--decisions", dec3, "--out", out, "--script", "el"])),
               SystemExit, "not in the adversarial file"))
    nore = dict(full); nore["2"] = {"decision": "accept"}
    dec4 = os.path.join(d, "d4.json")
    with open(dec4, "w") as fh: json.dump(nore, fh)
    ok("reconcile refuses a decision without a reason",
       refuses(lambda: reconcile(parse_args(["--export", e1, "--primary", wbp, "--adversarial", advp,
                                             "--decisions", dec4, "--out", out, "--script", "el"])),
               SystemExit, "needs a reason"))

    inplace = os.path.join(d, "QA_inplace.xlsx")
    import shutil as _sh; _sh.copy2(wbp, inplace)
    reconcile(parse_args(["--export", e1, "--primary", inplace, "--adversarial", advp, "--decisions", dec,
                          "--out", inplace, "--script", "el"]))
    ok("reconciling in place keeps the pre-reconciliation workbook",
       os.path.exists(inplace.replace(".xlsx", ".pre-adv.xlsx")))

    # ── the entry point answers at all ──────────────────────────────────────────────────────
    stqa = os.path.join(HERE, "stqa.sh")
    for args in (["integrity", e1, "--script", "el", "--json"], ["fonts", "--list"], ["--help"]):
        r = subprocess.run(["bash", stqa] + args, capture_output=True, text=True)
        ok(f"stqa.sh {args[0]}", r.returncode == 0, r.stderr[-300:])

print(f"\n{passed} checks held")
