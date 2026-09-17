#!/usr/bin/env python3
"""Reconcile an adversarial file into the corrections workbook.

    stqa.sh reconcile --export export.xlsx --primary QA_R2.xlsx \\
        --adversarial QA_Adversarial_R2_X.xlsx --decisions decisions.json \\
        --out QA_R2.xlsx --script el --adversary "Fable 5.1 Max"

decisions.json — one entry per DROP / DISPUTE / ADD row of the adversarial file, keyed by that
file's own `#` (verdict-claim rows included):

    {"1": {"decision": "accept", "reason": "simulate shows NBSP introduced"},
     "3": {"decision": "reject", "reason": "source says 'use'; 'visit' is a meaning change"}}

Every non-CONFIRM row must have a decision and every decision must name a row that exists — the
script refuses both ways, so neither a silent discard nor a typo that quietly drops a decision
can reach the client.

Effects on the corrections workbook:
  DISPUTE accepted → Proposed Correction replaced by Adversary Proposed and/or Severity replaced;
                     Comment gets "ADV (<model>): accepted — <reason>"
  DISPUTE rejected → Comment gets "ADV (<model>): rejected — <reason>"
  DROP accepted    → row removed (listed in the summary);  DROP rejected → "… rejected — …"
  ADD accepted     → new row (severity, value, argument as Comment); ADD rejected → not added, listed
  CONFIRM          → Comment gets "ADV (<model>): confirmed"
  verdict (Ref V)  → reported in the summary only
The result passes the same column contract, and the same font check, as a freshly built workbook.
"""
import argparse
import json
import re
import shutil
import sys

from stqa_adversarial import ADV_COLS, load_primary
from stqa_workbook import QABook, check_rows, write_rows, _s


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    for a in ("--export", "--primary", "--adversarial", "--decisions", "--out"):
        ap.add_argument(a, required=True)
    ap.add_argument("--script", default="latin")
    ap.add_argument("--adversary", default="adversary")
    ap.add_argument("--font", default=None, help="house font; overrides the per-script choice")
    return ap.parse_args(argv)


def load_adversarial(path):
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    hdr = [_s(c).strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    if len(hdr) < 12 or hdr[0] != "#" or hdr[1] != "verdict":
        sys.exit(f"{path}: not an adversarial file — expected the 12-column schema, got {hdr[:3]}")
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(x is not None for x in r): continue
        rows.append(dict(zip(ADV_COLS, [_s(x) for x in (list(r) + [None] * 12)[:12]], strict=True)))
    return rows


def reconcile(a):
    book = QABook(a.export, script=a.script, font=a.font)
    byn = load_primary(a.primary)
    A = load_adversarial(a.adversarial)
    # Every other input error in this file exits with an operator-facing message; a missing or
    # malformed decisions file used to be the one that came out as a raw traceback.
    try:
        with open(a.decisions) as fh:
            D = json.load(fh)
    except OSError as e:
        sys.exit(f"reconcile: cannot read --decisions {a.decisions}: {e.strerror}")
    except json.JSONDecodeError as e:
        sys.exit(f"reconcile: --decisions {a.decisions} is not valid JSON (line {e.lineno}, column {e.colno}): {e.msg}")
    if not isinstance(D, dict):
        sys.exit(f"reconcile: --decisions {a.decisions} must hold a JSON object keyed by row number, not {type(D).__name__}")

    # the adversarial file covers every primary row exactly once — enforced at write time, re-checked
    # here because a hand-edited file is the way that promise gets lost
    refs = [r["ref"] for r in A if r["ref"] not in ("", "V")]
    unknown_ref = sorted(set(refs) - set(byn))
    missing_ref = sorted(set(byn) - set(refs), key=lambda x: int(x) if x.isdigit() else 0)
    if unknown_ref: sys.exit(f"adversarial file references rows the primary does not have: {unknown_ref}")
    if missing_ref: sys.exit(f"primary rows with no verdict in the adversarial file: {missing_ref}")
    if len(refs) != len(set(refs)):
        sys.exit(f"primary rows with more than one verdict: {sorted({r for r in refs if refs.count(r) > 1})}")
    need = {r["n"] for r in A if r["verdict"] != "CONFIRM"}
    if unknown := sorted(set(D) - {r["n"] for r in A}):
        sys.exit(f"decisions name rows that are not in the adversarial file: {unknown}")
    if no_dec := sorted(need - set(D), key=lambda x: int(x) if x.isdigit() else 0):
        sys.exit(f"non-CONFIRM adversarial rows without a decision: {no_dec} "
                 "— accept or reject each, with a reason")
    if extra := sorted(set(D) & {r["n"] for r in A if r["verdict"] == "CONFIRM"}):
        sys.exit(f"decisions for CONFIRM rows (nothing to decide): {extra}")

    dropped, added, rejected, verdicts = [], [], [], []
    tag = f"ADV ({a.adversary})"
    for r in A:
        if r["verdict"] == "CONFIRM":
            byn[r["ref"]]["com"] += f"\n{tag}: confirmed"
            continue
        d = D[r["n"]]
        if not isinstance(d, dict):     # a hand-edited decisions file: still an operator error,
            sys.exit(f"row {r['n']}: each decision must be an object with 'decision' and 'reason', "
                     f"not {type(d).__name__}")
        dec, why = d.get("decision"), _s(d.get("reason")).strip()
        if dec not in ("accept", "reject"):
            sys.exit(f"row {r['n']}: decision must be 'accept' or 'reject', got {dec!r}")
        if not why:
            sys.exit(f"row {r['n']}: a decision needs a reason that cites evidence")
        if r["ref"] == "V":
            verdicts.append((dec, r["prim"], r["arg"][:120], why)); continue
        if r["verdict"] == "DROP":
            if dec == "accept":
                dropped.append((r["ref"], byn[r["ref"]]["ev"] or byn[r["ref"]]["gid"], why))
                byn.pop(r["ref"])
            else:
                byn[r["ref"]]["com"] += f"\n{tag}: DROP rejected — {why}"
        elif r["verdict"] == "DISPUTE":
            p = byn[r["ref"]]
            if dec == "accept":
                if r["prop"]: p["prop"] = r["prop"]
                if r["sev"] and r["sev"] != p["sev"]:
                    p["com"] += f"\n{tag}: severity {p['sev']} → {r['sev']}"; p["sev"] = r["sev"]
                p["com"] += f"\n{tag}: accepted — {why}"
            else:
                p["com"] += f"\n{tag}: rejected — {why}"; rejected.append((f"#{r['ref']}", why))
        elif r["verdict"] == "ADD":
            if dec == "accept":
                cat = re.match(r"^\[(.*?)\]", r["arg"])
                com = r["arg"] + f"\n{tag}: added — {why}"
                if not r["prop"]:
                    com = re.sub(r"^\[.*?\] ", "", com)      # an empty value needs SOURCE→/POLICY:/… first
                byn[f"add-{r['n']}"] = dict(
                    n="", sev=r["sev"], cat=cat.group(1) if cat else "Adversarial add",
                    gid=r["gid"], ev=r["ev"], ql=r["ql"], src=r["src"], cur=r["cur"],
                    prop=r["prop"], com=com)
                added.append((r["ev"] or r["gid"], r["sev"]))
            else:
                rejected.append((f"ADD {r['ev'] or r['gid']}", why))
        else:
            sys.exit(f"row {r['n']}: unknown verdict {r['verdict']!r}")

    F = list(byn.values())
    check_rows(F, book.keys, book.script_re)
    if a.out == a.primary:
        # the adversarial file's Ref # numbers point at THIS file; overwriting it in place makes a
        # second reconcile run silently mis-target, so the pre-reconciliation state is kept
        backup = re.sub(r"\.xlsx$", "", a.primary) + ".pre-adv.xlsx"
        shutil.copy2(a.primary, backup)
        print(f"kept the pre-reconciliation workbook as {backup}", file=sys.stderr)
    counts, n = write_rows(F, a.out, rtl=book.rtl, script=book.script, font=a.font)
    return dict(out=a.out, counts=counts, rows=n, dropped=dropped, added=added,
                rejected=rejected, verdict_claims=verdicts)


def main():
    a = parse_args()
    res = reconcile(a)
    print(f"written {res['out']}: {res['counts']} ({res['rows']} rows)")
    for k in ("dropped", "added", "rejected", "verdict_claims"):
        print(f"{k}: {res[k]}")


if __name__ == "__main__":
    main()
