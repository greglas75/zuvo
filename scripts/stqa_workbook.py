#!/usr/bin/env python3
"""Corrections-workbook builder with the column contract enforced.

    from build_workbook import QABook
    qa = QABook("export.xlsx", script="el")            # script: key of qa_checks.SCRIPTS; rtl=True for ar/he
    qa.add("HIGH", "Orthography", "p1-q3", proposed="Πού κατοικείτε;",
           gloss='EN: "where (relative)" → "Where?" — interrogative needs the accent',
           comment="Introduced in R4.")
    qa.add("MEDIUM", "Grid identity", "p3-q9-a39", repl=[("ιδιώτη\u00a0bookmaker", "ιδιώτη bookmaker")],
           gloss="EN: same text; NBSP → space", comment="…")
    qa.flag("Source defect", "p3-q10-a52", "", "SOURCE→ Sports bets … private bookmaker\\nStray <br> …")
    qa.flag("Hidden CSS", "p3-q11-a54", qa.source("p3-q11-a54"), "DESIGN: …")   # value == source is allowed
    qa.write("PROJECT_EL_QA_Corrections_R2.xlsx")       # next to the export, unless the user names a place

Rows are addressed by `Export Variable`, or by `<Global Id>::<Export Variable>` when the EV is
empty or repeated — the Translation File exporter emits `…|0` rows with no EV, and those are
exactly the rows a hidden-CSS or `<br />`-label finding lands on.

write() asserts the contract (references/workbook-contract.md) and refuses to save otherwise —
the font included: every cell is written in a face that renders the target script (fonts.py).
"""
import re
import sys
from collections import Counter

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:
    sys.exit("openpyxl is missing — run through scripts/stqa.sh (it bootstraps a venv), "
             "or: python3 -m pip install openpyxl")

import stqa_fonts as fonts

NB = "\u00a0"
NBSP_MARK = "{NBSP}"   # NOT U+237D \u237d: no Office font has that glyph, so it shipped as a tofu box
ORDER = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3, "FLAG": 4}
FILLS = {"CRITICAL": "C00000", "HIGH": "ED7D31", "MEDIUM": "FFC000", "LOW": "A9D08E", "FLAG": "8EA9DB"}
HDR = ["#", "Severity", "Category", "Global Id", "Export Variable", "Q Label", "Original Text (EN)",
       "Current Translation", "Proposed Correction", "Comment"]
QA_COLS = ["n", "sev", "cat", "gid", "ev", "ql", "src", "cur", "prop", "com"]
META = (r"\b(master|source|target|col(umn)?s?|revert|keep|remove|strip|mirror|see|comment"
        r"|global|script|apply|n/a|tbd)\b")
PREFIX = r"^(SOURCE→|POLICY:|DESIGN:|INFO:)"
HEDGES = r"\b(maybe|perhaps|possibly|might|seems?|arguably|i think|probably)\b"
SCRIPTS = {"el": r"[Ͱ-Ͽἀ-῿]", "cyr": r"[Ѐ-ӿ]", "ar": r"[؀-ۿ]",
           "he": r"[֐-׿]", "hy": r"[԰-֏]", "ka": r"[Ⴀ-ჿ]", "th": r"[฀-๿]",
           "ja": r"[぀-ヿ一-鿿]", "ko": r"[가-힯]", "zh": r"[一-鿿]", "latin": None}


def _s(x): return "" if x is None else str(x)


def mark(x): return _s(x).replace(NB, NBSP_MARK)


def frag(a, b):
    """Minimal changed fragments, widened to word boundaries."""
    i = 0
    while i < min(len(a), len(b)) and a[i] == b[i]: i += 1
    j = 0
    while j < min(len(a), len(b)) - i and a[-1 - j] == b[-1 - j]: j += 1
    while i > 0 and not a[i - 1].isspace(): i -= 1
    while j > 0 and not a[len(a) - j].isspace(): j -= 1
    return a[i:len(a) - j].strip(), b[i:len(b) - j].strip()


def load_export(path):
    """Export rows, with the parent Question Label filled in. Fails loudly on a file that is not
    an export: a corrections workbook passed by mistake would otherwise build a confident nonsense."""
    ws = load_workbook(path).active
    hdr = [_s(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    low = [h.strip().lower() for h in hdr]
    if len(low) < 5 or "global id" not in low[0] or "export variable" not in low[1]:
        sys.exit(f"{path}: not a translation export — expected 'Global Id | Export Variable | "
                 f"Question Label | Original Text | Translated Text', got {hdr[:5]}")
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        r = list(r) + [None] * 7
        rows.append(dict(gid=_s(r[0]), ev=_s(r[1]), ql=_s(r[2]), src=_s(r[3]), tgt=_s(r[4])))
    cur = None
    for r in rows:
        if r["ql"]: cur = r["ql"]
        r["qlabel"] = r["ql"] or cur
    for r in rows:                                    # statements 0|qid|sid → parent question label
        if r["gid"].startswith("0|"):
            parent = r["gid"].split("|")[1]
            for q in rows:
                if q["ql"] and q["gid"].count("|") == 1 and q["gid"].endswith("|" + parent):
                    r["qlabel"] = q["ql"]; break
    return rows


class QABook:
    def __init__(self, export_path, script="latin", rtl=False, font=None):
        if script not in SCRIPTS:
            sys.exit(f"unknown script {script!r}; choose one of {', '.join(SCRIPTS)}")
        self.script, self.script_re = script, SCRIPTS[script]
        self.rtl = rtl or script in ("ar", "he")
        self.font = font
        self.rows = load_export(export_path)
        self.by, self._dupes = {}, set()
        for r in self.rows:
            if not r["ev"]: continue
            if r["ev"] in self.by: self._dupes.add(r["ev"])
            self.by[r["ev"]] = r
        self.by_key = {f'{r["gid"]}::{r["ev"]}': r for r in self.rows}
        self.keys = {(r["gid"], r["ev"]) for r in self.rows}
        self.F = []

    def row(self, addr):
        """Export row by `Export Variable`, or by `<Global Id>::<Export Variable>`."""
        addr = _s(addr)
        if "::" in addr:
            if addr in self.by_key: return self.by_key[addr]
            raise KeyError(f"{addr}: no such (Global Id, Export Variable) pair in the export")
        if addr in self._dupes:
            raise KeyError(f"{addr!r} occurs more than once in the export "
                           f"— address it as <Global Id>::{addr}")
        if addr in self.by: return self.by[addr]
        raise KeyError(f"{addr!r} is not an Export Variable of this export"
                       + ("" if addr else "; an empty Export Variable is addressed as <Global Id>::"))

    def current(self, addr): return self.row(addr)["tgt"]
    def source(self, addr): return self.row(addr)["src"]

    def add(self, sev, cat, addr, proposed=None, repl=None, gloss="", comment=""):
        """A row with a target value: either `proposed` (the whole cell) or `repl` [(old, new), …]
        applied to the current cell. A value change needs an EN gloss — the reader does not read the
        target language and cannot spot five changed words in a 150-character cell."""
        r = self.row(addr); cur = r["tgt"]
        if repl:
            new = cur
            for a, b in repl:
                assert a in new, f"{addr}: {a!r} not in the current text"
                new = new.replace(a, b)
            proposed = new
        assert proposed is not None, addr
        if proposed and proposed != cur:
            oa, ob = frag(cur, proposed)
            comment = f'CHANGE: "{mark(oa)}" → "{mark(ob)}"\n' + (gloss + "\n" if gloss else "") + comment
        self.F.append(dict(sev=sev, cat=cat, gid=r["gid"], ev=r["ev"], ql=r["qlabel"], src=r["src"],
                           cur=cur, prop=proposed, com=comment))

    def flag(self, cat, addr, proposed, comment):
        """FLAG row: `proposed` is "" or an exact alternative value (a value equal to the source is fine)."""
        r = self.row(addr)
        if proposed and proposed != r["tgt"]:
            shown = (f'"(= Original Text (EN) verbatim, {len(proposed)} chars)"'
                     if proposed == r["src"] else '"' + mark(proposed[:80]) + '"')
            comment = f'CHANGE: "{mark(r["tgt"][:80])}" → {shown}\n' + comment
        self.F.append(dict(sev="FLAG", cat=cat, gid=r["gid"], ev=r["ev"], ql=r["qlabel"], src=r["src"],
                           cur=r["tgt"], prop=proposed, com=comment))

    def check(self):
        check_rows(self.F, self.keys, self.script_re)

    def write(self, out_path, title="QA Corrections"):
        self.check()
        return write_rows(self.F, out_path, rtl=self.rtl, title=title, script=self.script, font=self.font)


def value_contract(v, src, cur, script_re, where):
    """The rule for any cell an apply-script writes verbatim (Proposed Correction, Adversary
    Proposed): an exact target-cell value, or empty. A value byte-equal to the source is always fine."""
    v, src, cur = _s(v), _s(src), _s(cur)
    assert not re.fullmatch(r"[\s—–\-]+", v), f"placeholder instead of a value: {where}"
    if v and v != src:
        assert not re.search(r"→|" + META, v, re.I), f"meta-vocabulary in a value cell: {where}"
        if script_re and re.search(script_re, cur):
            assert re.search(script_re, v), f"value not in the target script: {where}"
        assert NB not in v or (NB in src and "{{" in v), f"NBSP in a value cell: {where}"


def check_rows(F, keys, script_re):
    """Column contract (references/workbook-contract.md). Raises AssertionError on the first breach."""
    for f in F:
        v, where = _s(f["prop"]), (f["ev"] or f["gid"])
        assert (f["gid"], f["ev"]) in keys, f"synthetic row: {where}"
        assert f["sev"] in ORDER, f"bad severity: {where}"
        value_contract(v, f["src"], f["cur"], script_re, where)
        if v == "":
            assert re.match(PREFIX, f["com"]), \
                f"empty Proposed Correction without a SOURCE→/POLICY:/DESIGN:/INFO: comment: {where}"
        else:
            assert v != f["cur"] or f["sev"] == "FLAG", f"no-op proposal: {where}"
            if v != f["cur"]:
                assert f["com"].startswith("CHANGE"), f"comment without a CHANGE line: {where}"
                assert re.search(r"(?m)^EN:", f["com"]), f"value change without an EN gloss: {where}"
        if f["com"].startswith("SOURCE→"):
            assert f["com"].split("\n")[0][len("SOURCE→"):].strip(), f"SOURCE→ without a value: {where}"


def write_rows(F, out_path, rtl=False, title="QA Corrections", script="latin", font=None):
    F.sort(key=lambda f: ORDER[f["sev"]])
    wb = Workbook(); o = wb.active; o.title = title[:31]
    o.append(HDR)
    for i, f in enumerate(F, 1):
        o.append([i, f["sev"], f["cat"], f["gid"], f["ev"], f["ql"], f["src"], f["cur"], f["prop"], f["com"]])
    finish(wb, o, script=script, font=font, rtl=rtl, sev_col=1, fills=FILLS,
           widths=[5, 11, 34, 20, 16, 9, 48, 52, 52, 70], rtl_cols=(7, 8))
    wb.save(out_path)
    return dict(Counter(f["sev"] for f in F)), len(F)


def finish(wb, o, script, font, rtl, sev_col, fills, widths, rtl_cols, white_bold=("CRITICAL", "DROP")):
    """Style the sheet, choose the font from the text actually written, and refuse to save a
    workbook whose cells are not all in it. Returns the note to show the human."""
    texts = [_s(c.value) for row in o.iter_rows() for c in row]
    name, size, note = fonts.resolve(script, texts=texts, font=font)
    style_sheet(o, rtl, sev_col, fills, widths, rtl_cols, name, size, white_bold)
    fonts.set_default(wb, name, size)
    fonts.assert_workbook(wb, name, size)
    print(f"font: {note}", file=sys.stderr)
    return note


def style_sheet(o, rtl, sev_col, fills, widths, rtl_cols, font_name, font_size,
                white_bold=("CRITICAL", "DROP")):
    thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in o[1]:
        c.font = Font(name=font_name, size=font_size, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = border
    for row in o.iter_rows(min_row=2, max_row=o.max_row):
        for c in row:
            c.font = Font(name=font_name, size=font_size)
            c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = border
        if rtl:
            for i in rtl_cols:
                row[i].alignment = Alignment(wrap_text=True, vertical="top",
                                             horizontal="right", readingOrder=2)
        key = row[sev_col].value
        if key in fills:
            row[sev_col].fill = PatternFill("solid", fgColor=fills[key])
            row[sev_col].font = Font(name=font_name, size=font_size, bold=True,
                                     color=("FFFFFF" if key in white_bold else "000000"))
    for i, w in enumerate(widths, 1): o.column_dimensions[get_column_letter(i)].width = w
    o.freeze_panes = "A2"; o.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{o.max_row}"
    # the sheet itself reads right-to-left, not only the cells
    if rtl: o.sheet_view.rightToLeft = True


if __name__ == "__main__":
    print(__doc__)
