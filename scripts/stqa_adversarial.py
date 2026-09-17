#!/usr/bin/env python3
"""Adversarial-review workbook builder (references/adversarial-review.md).

    from build_adversarial import AdvBook
    adv = AdvBook("export.xlsx", "PROJECT_EL_QA_Corrections_R2.xlsx", script="el",
                  primary_model="GPT Astra", adversary_model="Fable 5.1 Max")
    adv.confirm(1, "Reproduced: NBSP 22→1 after simulate; proposal grammatical; matches the termbase.")
    adv.dispute(7, severity="HIGH",
                argument="Q6x3 stem rephrases the construct; stem divergence in a battery is HIGH.")
    adv.dispute(12, proposed="Παίζω περίπου εξίσου σε φυσικούς χώρους και online",
                argument="…evidence…", gloss="EN: …")
    adv.drop(19, "Proposal introduces NBSP (simulate: NEW_issues_introduced_by_workbook lists this cell).")
    adv.add("MEDIUM", "Grid identity", "p3-q10-a51", proposed="…", gloss="EN: …",
            argument="a26/a38 got 'σε', a51 did not.")
    adv.verdict_dispute("GO", "HIGH row #4 is still open — GO is not supported.")
    adv.write("PROJECT_EL_QA_Adversarial_R2_Fable-5.1-Max.xlsx")

write() asserts: every primary row has exactly one verdict (CONFIRM included), every value obeys
the column contract, every Argument is evidence rather than a hedge, and every cell is in a font
that renders the target script. Sort DROP → DISPUTE → ADD → CONFIRM.
"""
import re
import sys
from collections import Counter

try:
    from openpyxl import load_workbook, Workbook
except ModuleNotFoundError:
    sys.exit("openpyxl is missing — run through scripts/stqa.sh (it bootstraps a venv), "
             "or: python3 -m pip install openpyxl")

from stqa_workbook import (QABook, HEDGES, ORDER, QA_COLS, finish, frag, mark, value_contract, _s)

ADV_HDR = ["#", "Verdict", "Ref #", "Severity (adversary)", "Global Id", "Export Variable", "Q Label",
           "Original Text (EN)", "Current Translation", "Primary Proposed", "Adversary Proposed", "Argument"]
ADV_COLS = ["n", "verdict", "ref", "sev", "gid", "ev", "ql", "src", "cur", "prim", "prop", "arg"]
ADV_ORDER = {"DROP": 0, "DISPUTE": 1, "ADD": 2, "CONFIRM": 3}
ADV_FILLS = {"DROP": "C00000", "DISPUTE": "ED7D31", "ADD": "FFC000", "CONFIRM": "A9D08E"}


def load_primary(path):
    """Rows of a corrections workbook, keyed by its own `#`. Refuses anything else: handing the
    adversary the export instead of the workbook must not silently produce an all-CONFIRM file."""
    ws = load_workbook(path).active
    hdr = [_s(c).strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    if len(hdr) < 10 or hdr[0] != "#" or "severity" not in hdr[1] or "proposed" not in hdr[8]:
        sys.exit(f"{path}: not a corrections workbook — expected the 10-column contract, got {hdr[:4]}")
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(x is not None for x in r): continue
        q = dict(zip(QA_COLS, [_s(x) for x in (list(r) + [None] * 10)[:10]], strict=True))
        assert q["n"] not in out, f"{path}: two rows numbered {q['n']}"
        out[q["n"]] = q
    return out


def unquoted(text):
    """Text outside "quoted spans" — a hedge inside a quotation is the primary's wording, not the
    adversary's."""
    return re.sub(r'"[^"]*"', " ", _s(text))


class AdvBook:
    def __init__(self, export_path, primary_wb_path, script="latin", primary_model="",
                 adversary_model="", font=None):
        self.book = QABook(export_path, script=script, font=font)     # export access + the shared contract
        self.primary = load_primary(primary_wb_path)
        self.models = (primary_model, adversary_model)
        self.font = font
        self.rows = []

    # ── verdicts ────────────────────────────────────────────────────────────────────────────
    def _ref(self, ref):
        ref = _s(ref)
        if ref not in self.primary:
            have = sorted(self.primary, key=lambda x: int(x) if x.isdigit() else 0)
            raise KeyError(f"no row #{ref} in the primary workbook; it has #{have[0]}–#{have[-1]}" if have
                           else "the primary workbook has no rows")
        return ref

    def _row(self, verdict, ref, sev, q, adv_prop, argument):
        self.rows.append(dict(verdict=verdict, ref=ref, sev=sev, gid=q["gid"], ev=q["ev"], ql=q["ql"],
                              src=q["src"], cur=q["cur"], prim=q.get("prop", ""), prop=adv_prop,
                              arg=argument))

    def confirm(self, ref, argument="Reproduced and read against the source; the row stands."):
        ref = self._ref(ref); q = self.primary[ref]
        self._row("CONFIRM", ref, q["sev"], q, "", argument)

    def dispute(self, ref, argument, severity=None, proposed=None, gloss=""):
        """`severity`: the tier the adversary argues for. `proposed`: the exact replacement value
        (None when only the severity is disputed)."""
        ref = self._ref(ref); q = self.primary[ref]
        assert severity or proposed is not None, f"#{ref}: a dispute needs a severity or a value"
        if proposed and proposed != q["cur"] and gloss:
            oa, ob = frag(q["cur"], proposed)
            argument = f'CHANGE: "{mark(oa)}" → "{mark(ob)}"\n{gloss}\n{argument}'
        self._row("DISPUTE", ref, severity or q["sev"], q, proposed or "", argument)

    def drop(self, ref, argument):
        ref = self._ref(ref); q = self.primary[ref]
        self._row("DROP", ref, q["sev"], q, "", argument)

    def add(self, sev, cat, addr, argument, proposed=None, gloss="", repl=None):
        """A finding the primary missed. `addr`: Export Variable or `<Global Id>::<Export Variable>`."""
        r = self.book.row(addr); cur = r["tgt"]
        if repl:
            proposed = cur
            for a, b in repl:
                assert a in proposed, f"{addr}: {a!r} not in the current text"
                proposed = proposed.replace(a, b)
        head = ""
        if proposed and proposed != cur:
            oa, ob = frag(cur, proposed)
            head = f'CHANGE: "{mark(oa)}" → "{mark(ob)}"\n' + (gloss + "\n" if gloss else "")
        q = dict(gid=r["gid"], ev=r["ev"], ql=r["qlabel"], src=r["src"], cur=cur, prop="")
        self._row("ADD", "", sev, q, proposed or "", head + f"[{cat}] " + argument)

    def verdict_dispute(self, claim, argument):
        """A claim of the primary's chat verdict the artifacts do not support."""
        q = dict(gid="", ev="", ql="", src="", cur="", prop=claim)
        self._row("DISPUTE", "V", "", q, "", argument)

    # ── contract ────────────────────────────────────────────────────────────────────────────
    def check(self):
        refs = Counter(r["ref"] for r in self.rows if r["ref"] not in ("", "V"))
        missing = sorted((n for n in self.primary if refs[n] == 0),
                         key=lambda x: int(x) if x.isdigit() else 0)
        dup = [n for n, c in refs.items() if c > 1]
        assert not missing, f"primary rows without a verdict (CONFIRM is not optional): {missing}"
        assert not dup, f"primary rows with more than one verdict: {dup}"
        for r in self.rows:
            where = f"{r['verdict']} {r['ref'] or r['ev']}"
            assert r["verdict"] in ADV_ORDER, where
            assert len(r["arg"].strip()) >= 20, f"Argument too thin: {where}"
            assert not re.search(HEDGES, unquoted(r["arg"]), re.I), \
                f"hedging outside a quotation — state the evidence: {where}"
            value_contract(r["prop"], r["src"], r["cur"], self.book.script_re, where)
            if r["verdict"] == "ADD":
                assert (r["gid"], r["ev"]) in self.book.keys, f"ADD with a synthetic key: {where}"
                assert r["sev"] in ORDER, f"ADD without a valid severity: {where}"
                assert r["prop"] or re.match(r"^(\[.*?\] )?(SOURCE→|POLICY:|DESIGN:|INFO:)", r["arg"]), \
                    f"ADD without a value or a prefixed argument: {where}"
                if r["prop"] and r["prop"] != r["cur"]:
                    assert re.search(r"(?m)^EN:", r["arg"]), \
                        f"ADD changes a value without an EN gloss: {where}"
            if r["verdict"] == "DISPUTE" and r["ref"] != "V":
                assert r["prop"] or r["sev"] != self.primary[r["ref"]]["sev"], \
                    f"DISPUTE changes neither value nor severity — it repeats the primary: #{r['ref']}"
                if r["prop"] and r["prop"] == self.primary[r["ref"]]["prop"]:
                    assert r["sev"] != self.primary[r["ref"]]["sev"], \
                        f"DISPUTE proposes the primary's own value — it repeats the primary: #{r['ref']}"

    def write(self, out_path):
        self.check()
        self.rows.sort(key=lambda r: (ADV_ORDER[r["verdict"]], r["ref"] != "V",
                                      int(r["ref"]) if r["ref"].isdigit() else 0))
        wb = Workbook(); o = wb.active
        o.title = (f"Adversarial {self.models[1]}" if self.models[1] else "Adversarial")[:31]
        o.append(ADV_HDR)
        for i, r in enumerate(self.rows, 1):
            o.append([i, r["verdict"], r["ref"], r["sev"], r["gid"], r["ev"], r["ql"], r["src"], r["cur"],
                      r["prim"], r["prop"], r["arg"]])
        finish(wb, o, script=self.book.script, font=self.font, rtl=self.book.rtl, sev_col=1, fills=ADV_FILLS,
               widths=[5, 10, 6, 12, 20, 16, 9, 44, 46, 46, 46, 70], rtl_cols=(8, 9, 10))
        wb.save(out_path)
        return dict(Counter(r["verdict"] for r in self.rows)), len(self.rows)


if __name__ == "__main__":
    print(__doc__)
